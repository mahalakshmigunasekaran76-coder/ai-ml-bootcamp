"""
merge_all.py
------------
Merge Tool: finds every PowerPoint and Excel file in a folder and all the
folders nested inside it, then creates up to two files:

    Combined-PPT_2026-10-20_21-09-30.pptx     all the slides, in path order
    Combined-Excel_2026-10-20_21-09-30.xlsx   all the sheets, one per sheet

A file is only created when that type was found. Both names share the
date and time the merge started.

Large files: a created file bigger than SPLIT_OVER_MB (50 MB) is replaced
by parts of at most PART_MB (45 MB) each, numbered _01, _02, ...:

    Combined-PPT_2026-10-20_21-09-30_01.pptx
    Combined-PPT_2026-10-20_21-09-30_02.pptx

PowerPoint is split between slides; Excel is split between source files,
so the sheets of one source workbook stay together. A single slide or
workbook that is larger than PART_MB on its own gets a part to itself.

Usage:
    pip install python-pptx openpyxl
    python merge_all_gui.py                          (the window)
    python merge_all.py "C:\\path\\to\\folder" -o "C:\\path\\to\\save"

These files must sit in the same folder:
    merge_all.py           finds the files and runs both merges (this file)
    merge_all_gui.py       the window
    ppt_merge_engine.py    merges the slides  (same as the PPT Merge Tool)
    excel_merge_engine.py  merges the sheets  (same as the Excel Merge Tool)
"""

import argparse
import contextlib
import io
import shutil
import sys
import tempfile
import zlib
from datetime import datetime
from pathlib import Path

PPT_TYPES = {".pptx"}
EXCEL_TYPES = {".xlsx", ".xlsm", ".csv"}
OLD_TYPES = {".ppt": "old .ppt format, save it as .pptx",
             ".xls": "old .xls format, save it as .xlsx",
             ".xlsb": "binary .xlsb format, save it as .xlsx"}
OUTPUT_PREFIXES = ("combinedppt_", "combinedexcel_", "combined-ppt_", "combined-excel_")

# Size rule: files above SPLIT_OVER_MB are split into parts of up to PART_MB.
# MB here is the same unit File Explorer shows (1 MB = 1024 x 1024 bytes).
SPLIT_OVER_MB = 50
PART_MB = 45
MB = 1024 * 1024


def find_files(root: Path):
    """Return (ppt_files, excel_files) found anywhere under root,
    including root itself, sorted by their path inside root."""
    ppt_files, excel_files = [], []
    for f in sorted(root.rglob("*"), key=lambda p: str(p.relative_to(root)).lower()):
        if not f.is_file():
            continue
        name, ext = f.name.lower(), f.suffix.lower()
        rel = f.relative_to(root)
        if name.startswith("~$"):                    # Office lock file
            continue
        if name.startswith(OUTPUT_PREFIXES):         # a result of an earlier run
            print(f"  [skip]  {rel}: created by an earlier merge")
            continue
        if ext in PPT_TYPES:
            print(f"  [found] {rel}")
            ppt_files.append(f)
        elif ext in EXCEL_TYPES:
            print(f"  [found] {rel}")
            excel_files.append(f)
        elif ext in OLD_TYPES:
            print(f"  [skip]  {rel}: {OLD_TYPES[ext]}")
    return ppt_files, excel_files


def output_names(stamp: datetime = None):
    """File names for this run, e.g. Combined-PPT_2026-10-20_21-09-30.pptx.
    Colons are not allowed in Windows file names, so the time uses dashes."""
    text = (stamp or datetime.now()).strftime("%Y-%m-%d_%H-%M-%S")
    return f"Combined-PPT_{text}.pptx", f"Combined-Excel_{text}.xlsx"


# --------------------------------------------------------------------------- #
# Splitting large files
# --------------------------------------------------------------------------- #
def _pack(items, budget):
    """Group items in order so each group's estimated size stays within budget.
    Each item is {file name: size}; a file used by several items in the same
    group (such as one picture on two slides) is counted once."""
    groups, current, files = [], [], {}
    for i, item in enumerate(items):
        if current and sum({**files, **item}.values()) > budget:
            groups.append(current)
            current, files = [], {}
        current.append(i)
        files.update(item)
    if current:
        groups.append(current)
    return groups


def _build_parts(groups, save_group, folder: Path, ext: str):
    """Save each group with save_group(items, path). A part that comes out
    bigger than PART_MB is split in half and saved again, until it fits or
    holds a single item. Returns [(path, items)] in order."""
    done, queue, n = [], list(groups), 0
    while queue:
        group = queue.pop(0)
        n += 1
        path = folder / f"part{n}{ext}"
        if not save_group(group, path):
            continue
        if path.stat().st_size > PART_MB * MB and len(group) > 1:
            path.unlink()
            half = len(group) // 2
            queue[0:0] = [group[:half], group[half:]]     # keep the order
            continue
        done.append((path, group))
    return done


def _split_if_large(full: Path, target: Path, make_parts, unit: str) -> list:
    """full is the merged file, built in a temporary folder. If it is at most
    SPLIT_OVER_MB it is moved to target. Otherwise numbered parts made by
    make_parts(folder) are moved there instead, so the big file never lands
    in the Output Path. Returns the paths created."""
    size = full.stat().st_size
    if size <= SPLIT_OVER_MB * MB:
        shutil.move(str(full), str(target))
        return [target]
    print(f"  {target.name} would be {size / MB:.1f} MB, over {SPLIT_OVER_MB} MB: "
          f"splitting it into parts of up to {PART_MB} MB")
    created = []
    with tempfile.TemporaryDirectory() as tmp:
        try:
            parts = make_parts(Path(tmp))
        except Exception as e:                       # keep the full file rather than nothing
            print(f"     [note] {target.name} could not be split ({type(e).__name__}: {e}); "
                  "the full file was kept")
            shutil.move(str(full), str(target))
            return [target]
        for i, (path, items) in enumerate(parts, 1):
            final = target.with_name(f"{target.stem}_{i:02d}{target.suffix}")
            shutil.move(str(path), str(final))
            part_mb = final.stat().st_size / MB
            shown = f"{part_mb:.1f} MB" if part_mb >= 0.1 else f"{final.stat().st_size / 1024:.0f} KB"
            print(f"  -> {final.name} ({len(items)} {unit}{'' if len(items) == 1 else 's'}, {shown})")
            if part_mb > PART_MB:
                print(f"     [note] this {unit} alone is bigger than {PART_MB} MB, so this part could not be made smaller")
            created.append(final)
    return created


# PowerPoint: split between slides ------------------------------------------ #
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SECTION_SLIDE = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"
SECTION = "{http://schemas.microsoft.com/office/powerpoint/2010/main}section"
SHARED_PARTS = ("slideLayout", "slideMaster", "notesMaster", "handoutMaster", "theme")


def _slide_files(part, found=None) -> dict:
    """The files a slide uses (the slide itself, pictures, charts, video,
    notes) with their estimated zipped sizes, leaving out the design that
    all slides share."""
    found = {} if found is None else found
    if str(part.partname) in found:
        return found
    blob = part.blob
    found[str(part.partname)] = len(zlib.compress(blob, 1)) if part.content_type.endswith("xml") else len(blob)
    for rel in part.rels.values():
        if rel.is_external or rel.reltype.endswith("/slide") or any(k in rel.reltype for k in SHARED_PARTS):
            continue
        _slide_files(rel.target_part, found)
    return found


def _save_slides(source: Path, keep, dest: Path) -> bool:
    """Save a copy of source holding only the slides at the positions in keep."""
    from pptx import Presentation
    prs = Presentation(str(source))
    keep = set(keep)
    slide_list = prs.slides._sldIdLst
    pres = prs.part._element
    for i, slide_id in reversed(list(enumerate(list(slide_list)))):
        if i in keep:
            continue
        rid, number = slide_id.rId, slide_id.get("id")
        slide_list.remove(slide_id)
        for el in list(pres.iter(SECTION_SLIDE)):            # sections
            if el.get("id") == number:
                el.getparent().remove(el)
        for el in list(pres.iter(f"{{{PRESENTATION_NS}}}sld")):   # custom shows
            if el.get(f"{{{REL_NS}}}id") == rid:
                el.getparent().remove(el)
        prs.part.drop_rel(rid)                               # unused files are left out on save
    for section in list(pres.iter(SECTION)):                 # sections left empty
        if not any(True for _ in section.iter(SECTION_SLIDE)):
            section.getparent().remove(section)
    prs.save(str(dest))
    return True


def _split_pptx(merged: Path, folder: Path):
    from pptx import Presentation
    prs = Presentation(str(merged))
    slides = [_slide_files(slide.part) for slide in prs.slides]
    _save_slides(merged, [], folder / "design.pptx")         # the shared design alone
    design = (folder / "design.pptx").stat().st_size
    budget = max(PART_MB * MB - design, PART_MB * MB / 10)
    return _build_parts(_pack(slides, budget), lambda keep, path: _save_slides(merged, keep, path),
                        folder, ".pptx")


# Excel: split between source files ----------------------------------------- #
def _split_excel(files, root: Path, folder: Path):
    from excel_merge_engine import merge_excel_files

    def save(group, path):
        try:
            with contextlib.redirect_stdout(io.StringIO()):  # the full merge already reported
                merge_excel_files([files[i] for i in group], path, root=root)
            return True
        except SystemExit:
            return False

    sizes = [{str(f): f.stat().st_size} for f in files]
    return _build_parts(_pack(sizes, PART_MB * MB), save, folder, ".xlsx")


# --------------------------------------------------------------------------- #
# Merging
# --------------------------------------------------------------------------- #
def merge_all(input_dir: Path, output_dir: Path) -> list:
    """Find the files, run each merge, and return the paths created.
    Raises ValueError with a readable message when nothing can be made."""
    input_dir, output_dir = Path(input_dir), Path(output_dir)
    if not input_dir.is_dir():
        raise ValueError(f"Folder not found:\n{input_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Scanning {input_dir} ...")
    ppt_files, excel_files = find_files(input_dir)
    if not ppt_files and not excel_files:
        raise ValueError("No PowerPoint (.pptx) or Excel (.xlsx, .xlsm, .csv) files "
                         f"were found in\n{input_dir}\nor the folders inside it.")

    ppt_name, excel_name = output_names()
    created, problems = [], []

    if ppt_files:
        print(f"\nPowerPoint: {len(ppt_files)} file(s)")
        from ppt_merge_engine import merge_with_python_pptx
        target = output_dir / ppt_name
        try:
            with tempfile.TemporaryDirectory() as work:      # build the full file outside the Output Path
                full = Path(work) / ppt_name
                slides = merge_with_python_pptx(ppt_files, full)
                parts = _split_if_large(full, target, lambda folder: _split_pptx(full, folder), "slide")
            if parts == [target]:
                print(f"  -> {ppt_name} ({slides} slides)")
            created.extend(parts)
        except SystemExit as e:              # the engine stops with a message
            problems.append(f"PowerPoint: {e}")
            print(f"  PowerPoint merge stopped: {e}")

    if excel_files:
        print(f"\nExcel: {len(excel_files)} file(s)")
        from excel_merge_engine import merge_excel_files
        target = output_dir / excel_name
        try:
            with tempfile.TemporaryDirectory() as work:
                full = Path(work) / excel_name
                sheets = merge_excel_files(excel_files, full, root=input_dir)
                parts = _split_if_large(full, target, lambda folder: _split_excel(excel_files, input_dir, folder), "file")
            if parts == [target]:
                print(f"  -> {excel_name} ({sheets} sheets)")
            created.extend(parts)
        except SystemExit as e:
            problems.append(f"Excel: {e}")
            print(f"  Excel merge stopped: {e}")

    if not created:
        raise ValueError("No files could be created.\n\n" + "\n".join(problems))
    return created


def main():
    parser = argparse.ArgumentParser(
        description="Merge all PowerPoint files into one .pptx and all Excel files "
                    "into one .xlsx, searching every nested folder. Files over "
                    f"{SPLIT_OVER_MB} MB are split into parts of up to {PART_MB} MB.")
    parser.add_argument("root", help="Folder that contains the files")
    parser.add_argument("-o", "--output-dir", default=None,
                        help="Folder to save the created files (default: the same folder)")
    args = parser.parse_args()
    try:
        created = merge_all(Path(args.root), Path(args.output_dir or args.root))
    except ValueError as e:
        sys.exit(str(e))
    print("\nFiles created:")
    for f in created:
        print(f"  {f}")


if __name__ == "__main__":
    main()
