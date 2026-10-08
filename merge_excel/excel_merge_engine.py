"""
excel_merge_engine.py
---------------------
Part 2 of the Excel Merge Tool: copies every sheet into one workbook.

Each sheet of each file becomes one sheet in the merged workbook:
  * a file with one sheet      -> sheet named after the file   ("Sales")
  * a file with several sheets -> "file - sheet"               ("Sales - Q1")
Names are shortened to Excel's 31-character limit and made unique.

Copied: values, formatting (fonts, fills, borders, number formats,
alignment), column widths, row heights, merged cells, frozen panes,
hyperlinks, comments, data validation, conditional formatting, images.
A "Contents" sheet at the front lists every sheet and where it came from.
"""

import copy
import csv
import re
import zipfile
from pathlib import Path

INVALID_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_NAME = 31


# --------------------------------------------------------------------------- #
# Sheet names
# --------------------------------------------------------------------------- #
def make_sheet_name(wanted: str, used: set) -> str:
    """Clean a name for Excel and make it unique (case-insensitive)."""
    base = INVALID_CHARS.sub("-", wanted).strip("' ") or "Sheet"
    name = base[:MAX_NAME]
    n = 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:MAX_NAME - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _quoted(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _rename_refs(formula: str, renames: dict) -> str:
    """Point formulas like =Data!A1 or ='My data'!A1 at the renamed sheets."""
    for old, new in renames.items():
        q_old = re.escape("'" + old.replace("'", "''") + "'!")
        formula = re.sub(q_old, _quoted(new) + "!", formula)
        if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", old):
            formula = re.sub(r"(?<![A-Za-z0-9_.'\]])" + re.escape(old) + "!",
                             _quoted(new) + "!", formula)
    return formula


# --------------------------------------------------------------------------- #
# Copying one sheet
# --------------------------------------------------------------------------- #
def _number(text: str):
    """Turn '42' or '3.5' from a CSV into a number; keep anything else as text."""
    t = text.strip()
    if re.fullmatch(r"-?\d+", t) and not (len(t) > 1 and t.lstrip("-").startswith("0")):
        return int(t)
    if re.fullmatch(r"-?\d*\.\d+", t):
        return float(t)
    return text


def copy_csv(path: Path, dest_ws):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, newline="", encoding=enc) as fh:
                rows = list(csv.reader(fh))
            break
        except UnicodeDecodeError:
            continue
    for row in rows:
        dest_ws.append([_number(v) for v in row])
    return len(rows), max((len(r) for r in rows), default=0)


def copy_sheet(src_ws, vals_ws, dest_ws, keep_formulas: bool, renames: dict):
    """Copy contents and formatting of src_ws into dest_ws.

    vals_ws is the same sheet loaded with saved results instead of formulas.
    """
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                saved = vals_ws[cell.coordinate].value if vals_ws is not None else None
                if keep_formulas or saved is None:
                    value = _rename_refs(value, renames)   # keep the formula
                else:
                    value = saved                          # keep the result
            elif value is not None and not isinstance(value, (str, int, float, bool)) \
                    and value.__class__.__name__ in ("ArrayFormula", "DataTableFormula"):
                value = vals_ws[cell.coordinate].value if vals_ws is not None else None
            new = dest_ws.cell(row=cell.row, column=cell.column, value=value)
            if cell.has_style:
                new.font = copy.copy(cell.font)
                new.fill = copy.copy(cell.fill)
                new.border = copy.copy(cell.border)
                new.alignment = copy.copy(cell.alignment)
                new.protection = copy.copy(cell.protection)
                new.number_format = cell.number_format
            if cell.hyperlink:
                new.hyperlink = copy.copy(cell.hyperlink)
            if cell.comment:
                new.comment = copy.copy(cell.comment)

    for rng in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(rng))
    for key, dim in src_ws.column_dimensions.items():
        d = dest_ws.column_dimensions[key]
        d.width, d.hidden = dim.width, dim.hidden
        d.min, d.max = dim.min, dim.max
    for key, dim in src_ws.row_dimensions.items():
        d = dest_ws.row_dimensions[key]
        d.height, d.hidden = dim.height, dim.hidden
    dest_ws.freeze_panes = src_ws.freeze_panes
    dest_ws.sheet_properties.tabColor = src_ws.sheet_properties.tabColor
    dest_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
    if src_ws.auto_filter.ref:
        dest_ws.auto_filter.ref = src_ws.auto_filter.ref

    for dv in src_ws.data_validations.dataValidation:
        dest_ws.add_data_validation(copy.copy(dv))
    for rng in src_ws.conditional_formatting:
        for rule in rng.rules:
            dest_ws.conditional_formatting.add(str(rng.sqref), copy.copy(rule))
    for img in getattr(src_ws, "_images", []):
        dest_ws.add_image(copy.copy(img))

    return src_ws.max_row, src_ws.max_column


def _has_parts(path: Path, folder: str) -> bool:
    try:
        with zipfile.ZipFile(path) as z:
            return any(n.startswith(folder) for n in z.namelist())
    except zipfile.BadZipFile:
        return False


# --------------------------------------------------------------------------- #
# Merging all files
# --------------------------------------------------------------------------- #
def merge_excel_files(files: list, output: Path, root: Path = None,
                      keep_formulas: bool = False, contents: bool = True) -> int:
    """Merge every sheet of every file into output. Returns the sheet count."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.hyperlink import Hyperlink
    except ImportError:
        raise SystemExit("openpyxl is not installed. Run: pip install openpyxl")

    merged = Workbook()
    toc = merged.active
    toc.title = "Contents"
    used = {"contents"}
    records = []                     # (sheet name, source, original sheet, rows, cols)

    for f in files:
        label = str(f.relative_to(root)) if root else f.name
        ext = f.suffix.lower()
        try:
            if ext == ".csv":
                name = make_sheet_name(f.stem, used)
                ws = merged.create_sheet(name)
                rows, cols = copy_csv(f, ws)
                records.append((name, label, "(CSV file)", rows, cols))
                print(f"  + {label}: 1 sheet added")
                continue

            src = load_workbook(f)                        # formulas
            vals = load_workbook(f, data_only=True)       # saved results
            sheets = [ws for ws in src.worksheets]
            if not sheets:
                print(f"  [skip] {label}: no worksheets")
                continue
            renames = {}
            for ws in sheets:
                wanted = f.stem if len(sheets) == 1 else f"{f.stem} - {ws.title}"
                renames[ws.title] = make_sheet_name(wanted, used)
            for ws in sheets:
                dest = merged.create_sheet(renames[ws.title])
                rows, cols = copy_sheet(ws, vals[ws.title], dest, keep_formulas, renames)
                records.append((dest.title, label, ws.title, rows, cols))
            print(f"  + {label}: {len(sheets)} sheet{'s' if len(sheets) > 1 else ''} added")

            if src.chartsheets or _has_parts(f, "xl/charts/"):
                print(f"    [note] {label} has charts; charts are not copied.")
            if _has_parts(f, "xl/pivotTables/"):
                print(f"    [note] {label} has pivot tables; their values are copied as plain cells.")
            if ext == ".xlsm":
                print(f"    [note] {label} has macros; macros are not copied.")
        except Exception as e:                            # damaged or protected file
            print(f"  [skip] {label}: could not be read ({type(e).__name__}: {e})")

    if not records:
        raise SystemExit("No sheets could be read from the files found.")

    if contents:
        bold = Font(name="Arial", bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", start_color="1F4E5F")
        toc.append(["#", "Sheet", "Source file", "Original sheet", "Rows", "Columns"])
        for c in toc[1]:
            c.font, c.fill = bold, head_fill
        for i, (name, label, orig, rows, cols) in enumerate(records, 1):
            toc.append([i, name, label, orig, rows, cols])
            link = toc.cell(row=i + 1, column=2)
            link.hyperlink = Hyperlink(ref=link.coordinate, location=f"{_quoted(name)}!A1")
            link.font = Font(name="Arial", color="0563C1", underline="single")
            for col in (1, 3, 4, 5, 6):
                toc.cell(row=i + 1, column=col).font = Font(name="Arial")
        for col, width in zip("ABCDEF", (5, 34, 50, 24, 9, 10)):
            toc.column_dimensions[col].width = width
        toc.freeze_panes = "A2"
        toc.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    else:
        merged.remove(toc)

    merged.save(output)
    return len(records)
