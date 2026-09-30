"""
merge_ppts.py
-------------
Collect PowerPoint files from every sub-folder of a root directory and merge
all of their slides into ONE presentation.

Example layout:
    my_decks/
        folder1/intro.pptx
        folder2/sales.pptx
        folder3/marketing.pptx
        folder4/finance.pptx
        folder5/summary.pptx

Usage:
    pip install python-pptx
    python merge_ppts.py my_decks -o combined.pptx

Options:
    -o / --output     Output file name (default: merged_presentation.pptx)
    --recursive       Also look in nested sub-folders (folder1/sub/...)
    --method com      Windows only: use installed PowerPoint (needs pywin32).
                      Handles .ppt files and complex content most faithfully.

Two merge engines:
  * "pptx" (default, any OS) - pure Python using python-pptx. Copies text,
    shapes, tables, pictures, charts, hyperlinks, slide backgrounds and
    speaker notes. The first deck is used as the base, so its theme/slide size
    is kept; slides from other decks are placed on the base's blank layout.
  * "com" (Windows + Microsoft PowerPoint) - asks PowerPoint itself to insert
    the slides, which also supports old .ppt files.
"""

import argparse
import copy
import re
import sys
from pathlib import Path

PPT_EXTENSIONS = {".pptx", ".ppt"}

# Relationship types we must NOT copy from a source slide: layout is replaced by
# the destination layout, and the others would drag in whole foreign slides.
SKIP_REL_KEYWORDS = ("slideLayout", "notesSlide", "/slide", "comments", "tags")

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# --------------------------------------------------------------------------- #
# 1. Find the files
# --------------------------------------------------------------------------- #
def find_ppt_files(root: Path, recursive: bool = False) -> list:
    """Return PowerPoint files from each sub-folder of root, in sorted order."""
    files = []
    subfolders = sorted(p for p in root.iterdir() if p.is_dir())
    if not subfolders:
        print(f"No sub-folders found in {root}")

    for folder in subfolders:
        pattern = "**/*" if recursive else "*"
        found = sorted(
            f for f in folder.glob(pattern)
            if f.is_file()
            and f.suffix.lower() in PPT_EXTENSIONS
            and not f.name.startswith("~$")      # skip PowerPoint lock files
        )
        if not found:
            print(f"  [skip] no PowerPoint file in {folder.name}")
        for f in found:
            print(f"  [found] {f.relative_to(root)}")
        files.extend(found)
    return files


# --------------------------------------------------------------------------- #
# 2a. Merge engine: python-pptx (cross-platform)
# --------------------------------------------------------------------------- #
def _should_skip(reltype: str) -> bool:
    if reltype.endswith("/slide"):           # link to another slide
        return True
    return any(k in reltype for k in SKIP_REL_KEYWORDS if k != "/slide")


def _partname_template(partname: str) -> str:
    """'/ppt/media/image3.png' -> '/ppt/media/image%d.png'"""
    m = re.match(r"^(.*?)(\d*)(\.[^./]+)$", partname)
    if not m:
        return partname + "%d"
    return f"{m.group(1)}%d{m.group(3)}"


def _adopt_part(part, package, renamed: set):
    """Give a foreign part (and its children) unique names in the new package,
    so e.g. two different 'image1.png' files don't overwrite each other."""
    if id(part) in renamed:
        return
    renamed.add(id(part))
    part.partname = package.next_partname(_partname_template(str(part.partname)))
    for rel in part.rels.values():
        if not rel.is_external:
            _adopt_part(rel.target_part, package, renamed)


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def copy_slide(src_slide, dest_prs, renamed: set):
    """Append a copy of src_slide to dest_prs."""
    package = dest_prs.part.package
    new_slide = dest_prs.slides.add_slide(_blank_layout(dest_prs))

    # Remove any placeholders the layout created
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)

    # Copy relationships (images, charts, media, hyperlinks) and map old->new ids
    rid_map = {}
    for rel in src_slide.part.rels.values():
        if _should_skip(rel.reltype):
            continue
        if rel.is_external:
            new_rid = new_slide.part.relate_to(rel.target_ref, rel.reltype,
                                               is_external=True)
        else:
            new_rid = new_slide.part.relate_to(rel.target_part, rel.reltype)
            _adopt_part(rel.target_part, package, renamed)
        rid_map[rel.rId] = new_rid

    def remap(element):
        for el in element.iter():
            for attr, val in list(el.attrib.items()):
                if attr.startswith("{%s}" % R_NS):
                    # unknown ids (e.g. skipped slide links) are dropped
                    if val in rid_map:
                        el.set(attr, rid_map[val])
                    else:
                        del el.attrib[attr]
        return element

    # Copy the slide's own background, if it has one
    src_bg = src_slide._element.cSld.bg
    if src_bg is not None:
        new_slide._element.cSld.insert(0, remap(copy.deepcopy(src_bg)))

    # Copy every shape (skip the two group-property elements at the start)
    dest_tree = new_slide.shapes._spTree
    for el in src_slide.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        dest_tree.append(remap(copy.deepcopy(el)))

    # Copy speaker notes (as plain text)
    if src_slide.has_notes_slide:
        notes = src_slide.notes_slide.notes_text_frame
        if notes is not None and notes.text.strip():
            new_slide.notes_slide.notes_text_frame.text = notes.text

    return new_slide


def merge_with_python_pptx(files: list, output: Path):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is not installed. Run: pip install python-pptx")

    pptx_files = [f for f in files if f.suffix.lower() == ".pptx"]
    for f in files:
        if f.suffix.lower() == ".ppt":
            print(f"  [warn] {f.name} is old .ppt format - skipped. "
                  "Re-save it as .pptx or use --method com.")
    if not pptx_files:
        sys.exit("No .pptx files to merge.")

    # First deck becomes the base (keeps its theme and slide size)
    merged = Presentation(str(pptx_files[0]))
    print(f"\nBase deck: {pptx_files[0].name} ({len(merged.slides)} slides)")

    renamed = set()
    for f in pptx_files[1:]:
        src = Presentation(str(f))
        if (src.slide_width, src.slide_height) != (merged.slide_width, merged.slide_height):
            print(f"  [note] {f.name} has a different slide size; "
                  "shapes keep their original positions.")
        for slide in src.slides:
            copy_slide(slide, merged, renamed)
        print(f"  + {f.name}: {len(src.slides)} slide(s) added")

    merged.save(str(output))
    return len(merged.slides)


# --------------------------------------------------------------------------- #
# 2b. Merge engine: Microsoft PowerPoint via COM (Windows only)
# --------------------------------------------------------------------------- #
def merge_with_powerpoint_com(files: list, output: Path):
    try:
        import win32com.client
    except ImportError:
        sys.exit("pywin32 is required for --method com. Run: pip install pywin32")

    app = win32com.client.Dispatch("PowerPoint.Application")
    merged = None
    try:
        merged = app.Presentations.Open(str(files[0].resolve()), WithWindow=False)
        print(f"\nBase deck: {files[0].name} ({merged.Slides.Count} slides)")
        for f in files[1:]:
            before = merged.Slides.Count
            merged.Slides.InsertFromFile(str(f.resolve()), before)
            print(f"  + {f.name}: {merged.Slides.Count - before} slide(s) added")
        merged.SaveAs(str(output.resolve()))
        return merged.Slides.Count
    finally:
        if merged is not None:
            merged.Close()
        app.Quit()


# --------------------------------------------------------------------------- #
# 3. Command line entry point
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(description="Merge PPT files from sub-folders.")
    parser.add_argument("root", help="Folder that contains the sub-folders")
    parser.add_argument("-o", "--output", default="merged_presentation.pptx",
                        help="Output file (default: merged_presentation.pptx)")
    parser.add_argument("--recursive", action="store_true",
                        help="Search nested sub-folders too")
    parser.add_argument("--method", choices=["pptx", "com"], default="pptx",
                        help="Merge engine (default: pptx)")
    args = parser.parse_args()

    root = Path(args.root)
    if not root.is_dir():
        sys.exit(f"Folder not found: {root}")

    output = Path(args.output)
    if output.suffix.lower() != ".pptx":
        output = output.with_suffix(".pptx")

    print(f"Scanning {root.resolve()} ...")
    files = [f for f in find_ppt_files(root, args.recursive)
             if f.resolve() != output.resolve()]
    if not files:
        sys.exit("No PowerPoint files found.")

    if args.method == "com":
        total = merge_with_powerpoint_com(files, output)
    else:
        total = merge_with_python_pptx(files, output)

    print(f"\nDone! {total} slides saved to {output.resolve()}")


if __name__ == "__main__":
    main()