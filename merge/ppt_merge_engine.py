"""
ppt_merge_engine.py
-------------------
Part 2 of the PPT Merge Tool: merges the slides into one presentation.

Two engines:
  * merge_with_python_pptx    - pure Python, works on any OS (default)
  * merge_with_powerpoint_com - uses installed PowerPoint (Windows only)
"""

import copy
import re
import sys
from pathlib import Path

# Relationship types we must NOT copy from a source slide: layout is replaced by
# the destination layout, and the others would drag in whole foreign slides.
SKIP_REL_KEYWORDS = ("slideLayout", "notesSlide", "/slide", "comments", "tags")

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------------------
# Engine 1: python-pptx (cross-platform)
# ---------------------------------------------------------------------------
def _should_skip(reltype: str) -> bool:
    if reltype.endswith("/slide"):           # link to another slide
        return True
    return any(k in reltype for k in SKIP_REL_KEYWORDS if k != "/slide")


def _partname_template(partname: str) -> str:
    """'/ppt/media/image3.png' -> '/ppt/media/image%d.png'"""
    m = re.match(r"^(.*?)(\d*)(\.[^./]+)$", partname)
    if not m:
        return partname + "%d"
    return f"{m.group(1)}%d{m.group(3)}"


def _adopt_part(part, package, renamed: set):
    """Give a foreign part (and its children) unique names in the new package,
    so e.g. two different 'image1.png' files don't overwrite each other."""
    if id(part) in renamed:
        return
    renamed.add(id(part))
    part.partname = package.next_partname(_partname_template(str(part.partname)))
    for rel in part.rels.values():
        if not rel.is_external:
            _adopt_part(rel.target_part, package, renamed)


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def _matching_layout(dest_prs, src_slide):
    """Find the base deck's layout with the same name as the source slide's
    layout (e.g. "Title Slide"), so titles and subtitles keep their position
    and formatting. Returns None if the base deck has no such layout."""
    name = src_slide.slide_layout.name
    for layout in dest_prs.slide_layouts:
        if layout.name == name:
            return layout
    return None


def _pin_position(src_shape, new_el):
    """For a placeholder with no position of its own, write the position it
    inherits from its layout onto the copy, so it doesn't jump to the corner."""
    if not src_shape.is_placeholder or new_el.xfrm is not None:
        return
    try:
        left, top = src_shape.left, src_shape.top
        width, height = src_shape.width, src_shape.height
    except (AttributeError, KeyError, ValueError):
        return
    if None in (left, top, width, height):
        return
    new_el.x, new_el.y, new_el.cx, new_el.cy = left, top, width, height


def copy_slide(src_slide, dest_prs, renamed: set):
    """Append a copy of src_slide to dest_prs."""
    package = dest_prs.part.package
    # Use the matching layout (e.g. "Title Slide") if the base deck has one;
    # otherwise fall back to the blank layout
    layout = _matching_layout(dest_prs, src_slide)
    new_slide = dest_prs.slides.add_slide(layout or _blank_layout(dest_prs))

    # Remove any placeholders the layout created
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)

    # Copy relationships (images, charts, media, hyperlinks) and map old->new ids
    rid_map = {}
    for rel in src_slide.part.rels.values():
        if _should_skip(rel.reltype):
            continue
        if rel.is_external:
            new_rid = new_slide.part.relate_to(rel.target_ref, rel.reltype,
                                               is_external=True)
        else:
            new_rid = new_slide.part.relate_to(rel.target_part, rel.reltype)
            _adopt_part(rel.target_part, package, renamed)
        rid_map[rel.rId] = new_rid

    def remap(element):
        for el in element.iter():
            for attr, val in list(el.attrib.items()):
                if attr.startswith("{%s}" % R_NS):
                    # unknown ids (e.g. skipped slide links) are dropped
                    if val in rid_map:
                        el.set(attr, rid_map[val])
                    else:
                        del el.attrib[attr]
        return element

    # Copy the slide's own background, if it has one
    src_bg = src_slide._element.cSld.bg
    if src_bg is not None:
        new_slide._element.cSld.insert(0, remap(copy.deepcopy(src_bg)))

    # Copy every shape (skip the two group-property elements at the start)
    src_shapes = {shape._element: shape for shape in src_slide.shapes}
    dest_tree = new_slide.shapes._spTree
    for el in src_slide.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        new_el = remap(copy.deepcopy(el))
        if layout is None and el in src_shapes:
            _pin_position(src_shapes[el], new_el)
        dest_tree.append(new_el)

    # Copy speaker notes (as plain text)
    if src_slide.has_notes_slide:
        notes = src_slide.notes_slide.notes_text_frame
        if notes is not None and notes.text.strip():
            new_slide.notes_slide.notes_text_frame.text = notes.text

    return new_slide


def merge_with_python_pptx(files: list, output: Path):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is not installed. Run: pip install python-pptx")

    pptx_files = [f for f in files if f.suffix.lower() == ".pptx"]
    for f in files:
        if f.suffix.lower() == ".ppt":
            print(f"  [warn] {f.name} is old .ppt format - skipped. "
                  "Re-save it as .pptx or use --method com.")
    if not pptx_files:
        sys.exit("No .pptx files to merge.")

    # First deck becomes the base (keeps its theme and slide size)
    merged = Presentation(str(pptx_files[0]))
    print(f"\nBase deck: {pptx_files[0].name} ({len(merged.slides)} slides)")

    renamed = set()
    for f in pptx_files[1:]:
        src = Presentation(str(f))
        if (src.slide_width, src.slide_height) != (merged.slide_width, merged.slide_height):
            print(f"  [note] {f.name} has a different slide size; "
                  "shapes keep their original positions.")
        for slide in src.slides:
            copy_slide(slide, merged, renamed)
        print(f"  + {f.name}: {len(src.slides)} slide(s) added")

    merged.save(str(output))
    return len(merged.slides)


# --------------------------------------------------------------------------- #
# Engine 2: Microsoft PowerPoint via COM (Windows only)
# --------------------------------------------------------------------------- #
def merge_with_powerpoint_com(files: list, output: Path):
    try:
        import win32com.client
    except ImportError:
        sys.exit("pywin32 is required for --method com. Run: pip install pywin32")

    app = win32com.client.Dispatch("PowerPoint.Application")
    merged = None
    try:
        merged = app.Presentations.Open(str(files[0].resolve()), WithWindow=False)
        print(f"\nBase deck: {files[0].name} ({merged.Slides.Count} slides)")
        for f in files[1:]:
            before = merged.Slides.Count
            merged.Slides.InsertFromFile(str(f.resolve()), before)
            print(f"  + {f.name}: {merged.Slides.Count - before} slide(s) added")
        merged.SaveAs(str(output.resolve()))
        return merged.Slides.Count
    finally:
        if merged is not None:
            merged.Close()
        app.Quit()
